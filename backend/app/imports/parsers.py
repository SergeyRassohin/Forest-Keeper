import csv
import json
import math
import posixpath
import re
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Protocol
from zipfile import BadZipFile, ZipFile

from defusedxml.ElementTree import iterparse, parse
from openpyxl import load_workbook


class ParseFailure(Exception):
    def __init__(self, code: str, message: str) -> None:
        self.code = code
        self.message = message


class XmlNode(Protocol):
    tag: str
    attrib: dict[str, str]
    text: str | None
    tail: str | None

    def __iter__(self) -> Iterator[Any]: ...


@dataclass
class ParsedRecord:
    position_key: str
    position: dict[str, Any]
    payload: dict[str, Any] = field(default_factory=dict)
    record_key: str | None = None
    object_type: str = "record"
    error_code: str | None = None
    error_message: str | None = None


def headers(values: list[Any]) -> list[str]:
    if not values or any(
        not isinstance(value, str) or not value or "\x00" in value for value in values
    ):
        raise ParseFailure("INVALID_HEADER", "Заголовки колонок должны быть непустыми строками")
    if len(set(values)) != len(values):
        raise ParseFailure("INVALID_HEADER", "Заголовки колонок не должны повторяться")
    return values


def json_safe(value: Any) -> bool:
    if isinstance(value, str):
        return "\x00" not in value
    if isinstance(value, float):
        return math.isfinite(value)
    if isinstance(value, dict):
        return all(json_safe(key) and json_safe(item) for key, item in value.items())
    if isinstance(value, list):
        return all(json_safe(item) for item in value)
    return True


def bounded(record: ParsedRecord) -> ParsedRecord:
    if not json_safe(record.payload) or len(record.object_type) > 200:
        record.payload = {}
        record.error_code = "INVALID_RAW_VALUE"
        record.error_message = "Запись нельзя сохранить без изменения исходных значений"
    elif (
        len(json.dumps([record.payload, record.position], ensure_ascii=False).encode()) > 1_000_000
    ):
        record.payload = {}
        # Keep a stable location even when workbook style/sheet metadata is excessive.
        record.position = {
            key: value
            for key, value in record.position.items()
            if key in {"row", "line_start", "line_end", "object"}
        }
        if record.position.get("row") is not None:
            record.position["sheet_index"] = 0
        record.error_code = "RECORD_TOO_LARGE"
        record.error_message = "Запись превышает допустимый размер 1 MB"
    return record


def csv_records(path: Path) -> Iterator[ParsedRecord]:
    csv.field_size_limit(1_000_000)
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream, strict=True)
        try:
            names = headers(next(reader))
        except StopIteration:
            raise ParseFailure("EMPTY_FILE", "Нет строки заголовков") from None
        number = 0
        while True:
            start = reader.line_num + 1
            try:
                row = next(reader)
            except StopIteration:
                break
            except csv.Error:
                number += 1
                yield ParsedRecord(
                    f"csv:{number}",
                    {"line_start": start, "line_end": reader.line_num},
                    error_code="INVALID_CSV_ROW",
                    error_message="Некорректная CSV-запись",
                )
                continue
            if not row:  # Empty physical lines are ignored; quoted empty cells are records.
                continue
            number += 1
            item = ParsedRecord(
                f"csv:{number}",
                {"line_start": start, "line_end": reader.line_num},
            )
            if len(row) != len(names):
                item.error_code = "COLUMN_COUNT"
                item.error_message = "Количество значений не совпадает с заголовком"
            else:
                item.payload = dict(zip(names, row, strict=True))
                item.record_key = item.payload.get("id")
            yield bounded(item)


def check_xlsx(path: Path, max_expanded_bytes: int) -> None:
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            if (
                len(members) > 10_000
                or sum(item.file_size for item in members) > max_expanded_bytes
            ):
                raise ParseFailure("UNSAFE_XLSX", "Превышен предел распаковки XLSX")
            if len({item.filename for item in members}) != len(members):
                raise ParseFailure("UNSAFE_XLSX", "Повторяющиеся элементы архива XLSX")
            for item in members:
                name = item.filename.replace("\\", "/")
                if (
                    name.startswith("/")
                    or ".." in name.split("/")
                    or item.flag_bits & 1
                    or "vbaproject" in name.lower()
                    or item.file_size > 200 * max(item.compress_size, 1)
                ):
                    raise ParseFailure("UNSAFE_XLSX", "Небезопасный архив XLSX")
            if "[Content_Types].xml" not in archive.namelist():
                raise ParseFailure("INVALID_XLSX", "Неверный формат XLSX")
            worksheets = {
                member.filename
                for member in members
                if member.filename.startswith("xl/worksheets/") and member.filename.endswith(".xml")
            }
            # OPC relationships may name worksheets outside this directory or with a
            # non-.xml extension. Check the actual parts that openpyxl will follow.
            for member in members:
                if not member.filename.endswith(".rels"):
                    continue
                if member.file_size > 1_000_000:
                    raise ParseFailure("UNSAFE_XLSX", "Слишком большие связи XLSX")
                with archive.open(member) as relationships:
                    root = parse(
                        relationships, forbid_dtd=True, forbid_entities=True, forbid_external=True
                    ).getroot()
                if root is None:
                    raise ParseFailure("INVALID_XLSX", "Отсутствуют связи XLSX")
                for relationship in root:
                    if not relationship.attrib.get("Type", "").endswith("/worksheet"):
                        continue
                    target = relationship.attrib.get("Target", "")
                    if relationship.attrib.get("TargetMode") == "External" or "\\" in target:
                        raise ParseFailure("UNSAFE_XLSX", "Недопустимая ссылка на лист XLSX")
                    directory = posixpath.dirname(posixpath.dirname(member.filename))
                    part = posixpath.normpath(
                        target.lstrip("/")
                        if target.startswith("/")
                        else posixpath.join(directory, target)
                    )
                    if part.startswith("../") or part not in archive.namelist():
                        raise ParseFailure("UNSAFE_XLSX", "Недопустимый путь листа XLSX")
                    worksheets.add(part)
            # Workbook sheet IDs are authoritative for what ExcelReader will open.
            # A forged relationship Type (e.g. image) must not hide a real worksheet.
            types = package_xml(archive, "[Content_Types].xml")
            main_types = {
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.template.main+xml",
            }
            workbooks = [
                entry.attrib.get("PartName", "").lstrip("/")
                for entry in types
                if entry.attrib.get("ContentType") in main_types
            ]
            if len(workbooks) != 1:
                raise ParseFailure("UNSAFE_XLSX", "Неподдерживаемая структура workbook XLSX")
            workbook_part = workbooks[0]
            workbook = package_xml(archive, workbook_part)
            workbook_rels = posixpath.join(
                posixpath.dirname(workbook_part),
                "_rels",
                posixpath.basename(workbook_part) + ".rels",
            )
            relationships = package_xml(archive, workbook_rels)
            by_id = {entry.attrib.get("Id"): entry.attrib for entry in relationships}
            if len(by_id) != len(relationships):
                raise ParseFailure("UNSAFE_XLSX", "Повторяющиеся связи workbook XLSX")
            prefix = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
            for sheet in workbook.iter():
                if sheet.tag.rsplit("}", 1)[-1] != "sheet":
                    continue
                sheet_relationship = by_id.get(sheet.attrib.get("{" + prefix[:-1] + "}id"))
                if (
                    sheet_relationship is None
                    or sheet_relationship.get("Type")
                    not in {
                        prefix + "worksheet",
                        prefix + "chartsheet",
                    }
                    or sheet_relationship.get("TargetMode") == "External"
                ):
                    raise ParseFailure("UNSAFE_XLSX", "Недопустимая связь листа workbook XLSX")
                target = sheet_relationship.get("Target", "")
                part = posixpath.normpath(
                    target.lstrip("/")
                    if target.startswith("/")
                    else posixpath.join(posixpath.dirname(workbook_part), target)
                )
                if part.startswith("../") or "\\" in part or part not in archive.namelist():
                    raise ParseFailure("UNSAFE_XLSX", "Недопустимый путь листа XLSX")
                if sheet_relationship["Type"] == prefix + "worksheet":
                    worksheets.add(part)
            for part in worksheets:
                with archive.open(part) as worksheet:
                    check_worksheet(worksheet)
    except BadZipFile:
        raise ParseFailure("INVALID_XLSX", "Повреждённый файл XLSX") from None


def package_xml(archive: ZipFile, name: str) -> Any:
    if name not in archive.namelist() or archive.getinfo(name).file_size > 1_000_000:
        raise ParseFailure("UNSAFE_XLSX", "Недопустимая metadata workbook XLSX")
    with archive.open(name) as stream:
        root = parse(stream, forbid_dtd=True, forbid_entities=True, forbid_external=True).getroot()
    if root is None:
        raise ParseFailure("INVALID_XLSX", "Отсутствует metadata workbook XLSX")
    return root


def check_worksheet(stream: Any) -> None:
    """Bound logical work too: a tiny ZIP can claim billions of empty Excel rows."""
    row, column = 0, 0
    stack: list[Any] = []
    for event, element in iterparse(
        stream,
        events=("start", "end"),
        forbid_dtd=True,
        forbid_entities=True,
        forbid_external=True,
    ):
        tag = element.tag.rsplit("}", 1)[-1]
        if event == "start":
            stack.append(element)
            if len(stack) > 64:
                raise ParseFailure("UNSAFE_XLSX", "Превышена вложенность XLSX")
            if tag == "row":
                raw_index = element.attrib.get("r", str(row + 1))
                if not raw_index.isdigit():
                    raise ParseFailure("UNSAFE_XLSX", "Недопустимая позиция строки XLSX")
                index = int(raw_index)
                if not row < index <= 1_048_576 or index - row > 1000:
                    raise ParseFailure(
                        "UNSAFE_XLSX", "Недопустимый или слишком большой разрыв строк XLSX"
                    )
                row, column = index, 0
            elif tag == "c":
                coordinate = element.attrib.get("r")
                if coordinate:
                    match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})", coordinate)
                    if match is None or int(match[2]) != row:
                        raise ParseFailure("UNSAFE_XLSX", "Недопустимая позиция ячейки XLSX")
                    index = 0
                    for letter in match[1]:
                        index = index * 26 + ord(letter) - ord("A") + 1
                else:
                    index = column + 1
                if not column < index <= 1024:
                    raise ParseFailure(
                        "UNSAFE_XLSX", "Превышен предел колонок XLSX или повторена ячейка"
                    )
                column = index
        else:
            stack.pop()
            if stack:
                stack[-1].remove(element)
            element.clear()


def json_cell(value: Any) -> Any:
    # JSON encoding of Excel's typed values is explicit, not business normalization.
    if isinstance(value, (datetime, date, time)):
        return {"iso8601": value.isoformat(), "excel_type": type(value).__name__}
    return value


def xlsx_records(path: Path, max_expanded_bytes: int) -> Iterator[ParsedRecord]:
    check_xlsx(path, max_expanded_bytes)
    with path.open("rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        try:
            if not workbook.worksheets:
                raise ParseFailure("INVALID_XLSX", "В XLSX нет листов данных")
            sheet = workbook.worksheets[0]
            # Do not trust declared dimensions from an untrusted workbook.
            sheet.reset_dimensions()
            rows = sheet.iter_rows()
            try:
                first = next(rows)
            except StopIteration:
                raise ParseFailure("INVALID_HEADER", "Нет строки заголовков") from None
            names = headers([cell.value for cell in first])
            for number, row in enumerate(rows, start=2):
                if all(cell.value is None for cell in row):
                    continue
                item = ParsedRecord(f"xlsx:{number}", {"sheet": sheet.title, "row": number})
                if len(row) > len(names) and any(
                    cell.value is not None for cell in row[len(names) :]
                ):
                    item.error_code = "COLUMN_COUNT"
                    item.error_message = "Значения выходят за пределы заголовков"
                else:
                    item.payload = {
                        name: json_cell(row[index].value) if index < len(row) else None
                        for index, name in enumerate(names)
                    }
                    item.position["cell_types"] = [cell.data_type for cell in row[: len(names)]]
                    item.position["number_formats"] = [
                        cell.number_format for cell in row[: len(names)]
                    ]
                    key = item.payload.get("id")
                    item.record_key = str(key) if key is not None else None
                yield bounded(item)
        finally:
            workbook.close()


def xml_payload(element: XmlNode, depth: int = 0) -> dict[str, Any]:
    if depth > 64:
        raise ParseFailure("XML_DEPTH", "Превышена допустимая вложенность XML")
    return {
        "tag": element.tag,
        "attributes": dict(element.attrib),
        "text": element.text,
        "tail": element.tail,
        "children": [xml_payload(child, depth + 1) for child in element],
    }


def xml_records(path: Path, one_c: bool) -> Iterator[ParsedRecord]:
    depth, number = 0, 0
    root: Any = None
    with path.open("rb") as stream:
        events = iterparse(
            stream,
            events=("start", "end"),
            forbid_dtd=True,
            forbid_entities=True,
            forbid_external=True,
        )
        for event, element in events:
            if event == "start":
                depth += 1
                if depth > 64:
                    raise ParseFailure("XML_DEPTH", "Превышена допустимая вложенность XML")
                if root is None:
                    root = element
                    if element.tag != ("Выгрузка" if one_c else "records"):
                        raise ParseFailure("XML_ROOT", "Неподдерживаемый корневой элемент XML")
            else:
                if depth == 2:
                    number += 1
                    item = ParsedRecord(
                        f"xml:{number}",
                        {"object": number, "tag": element.tag},
                        object_type=element.attrib.get("Тип", "record") if one_c else "record",
                    )
                    if element.tag != ("Объект" if one_c else "record") or (
                        one_c and (not element.attrib.get("Ид") or not element.attrib.get("Тип"))
                    ):
                        item.error_code = "INVALID_XML_OBJECT"
                        item.error_message = "Объект XML не соответствует формату выгрузки"
                    else:
                        item.payload = xml_payload(element)
                        item.record_key = element.attrib.get("Ид" if one_c else "id")
                    yield bounded(item)
                    if root is not None:
                        root.remove(element)
                    element.clear()
                depth -= 1


def parse_records(
    path: Path,
    file_format: str,
    source_type: str,
    max_expanded_bytes: int,
) -> Iterator[ParsedRecord]:
    try:
        if file_format == "CSV":
            yield from csv_records(path)
        elif file_format == "XLSX":
            yield from xlsx_records(path, max_expanded_bytes)
        elif file_format == "XML":
            yield from xml_records(path, source_type == "ONE_C")
        else:
            raise ParseFailure("UNSUPPORTED_FORMAT", "Неподдерживаемый формат")
    except ParseFailure:
        raise
    except Exception:
        # Library errors may include raw source data, paths or XML content.
        raise ParseFailure("INVALID_FILE", "Файл повреждён или не соответствует формату") from None
